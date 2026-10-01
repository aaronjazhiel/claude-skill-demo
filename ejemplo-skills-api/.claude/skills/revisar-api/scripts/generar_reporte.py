#!/usr/bin/env python3
"""Genera un Word desde datos JSON; no analiza el contrato de API."""
import argparse
import json
from pathlib import Path
from docx import Document

ROOT = Path(__file__).resolve().parents[1]


def validar(data):
    if not isinstance(data, dict):
        raise ValueError('El JSON debe contener un objeto.')
    for key in ('titulo', 'contrato', 'resumen'):
        if not isinstance(data.get(key), str) or not data[key].strip():
            raise ValueError(f'{key}: se requiere texto no vacío.')
    for key in ('hallazgos', 'preguntas'):
        if not isinstance(data.get(key), list):
            raise ValueError(f'{key}: se requiere una lista.')
    for i, item in enumerate(data['hallazgos']):
        if not isinstance(item, dict):
            raise ValueError(f'hallazgos[{i}]: se requiere un objeto.')
        for key in ('prioridad', 'tipo', 'hallazgo', 'evidencia', 'recomendacion'):
            if not isinstance(item.get(key), str) or not item[key].strip():
                raise ValueError(f'hallazgos[{i}].{key}: falta texto.')
        if item['prioridad'] not in ('Alta', 'Media', 'Baja'):
            raise ValueError(f'hallazgos[{i}].prioridad: usa Alta, Media o Baja.')
    if any(not isinstance(p, str) or not p.strip() for p in data['preguntas']):
        raise ValueError('Cada pregunta debe ser texto no vacío.')


def generar(datos, salida, sobrescribir=False):
    data = json.loads(datos.read_text(encoding='utf-8'))
    validar(data)
    if salida.exists() and not sobrescribir:
        raise ValueError('La salida ya existe. Elige otra ruta o usa --sobrescribir.')
    doc = Document(ROOT / 'templates' / 'reporte.docx')
    # Eliminar cuerpo de muestra preservando propiedades de sección y encabezado.
    body = doc._element.body
    for element in list(body):
        if element.tag.rsplit('}', 1)[-1] != 'sectPr':
            body.remove(element)
    doc.add_heading(data['titulo'], 0)
    doc.add_paragraph('Contrato: ' + data['contrato'])
    doc.add_heading('Resumen', 1)
    doc.add_paragraph(data['resumen'])
    doc.add_heading('Hallazgos', 1)
    if not data['hallazgos']:
        doc.add_paragraph('No se identificaron hallazgos en la revisión proporcionada.')
    for i, item in enumerate(data['hallazgos'], 1):
        doc.add_heading(f"{i}. {item['hallazgo']}", 2)
        for label, key in [('Prioridad', 'prioridad'), ('Tipo', 'tipo'),
                           ('Evidencia', 'evidencia'), ('Recomendación', 'recomendacion')]:
            p = doc.add_paragraph()
            p.add_run(label + ': ').bold = True
            p.add_run(item[key])
    doc.add_heading('Preguntas pendientes', 1)
    if not data['preguntas']:
        doc.add_paragraph('No se registraron preguntas pendientes.')
    for pregunta in data['preguntas']:
        doc.add_paragraph(pregunta, style='List Bullet')
    salida.parent.mkdir(parents=True, exist_ok=True)
    doc.save(salida)
    print('Reporte generado: ' + str(salida.resolve()))


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--datos', required=True, type=Path)
    parser.add_argument('--salida', required=True, type=Path)
    parser.add_argument('--sobrescribir', action='store_true')
    args = parser.parse_args()
    try:
        generar(args.datos, args.salida, args.sobrescribir)
    except (ValueError, OSError) as exc:
        parser.exit(2, f'Error: {exc}\n')


if __name__ == '__main__':
    main()
